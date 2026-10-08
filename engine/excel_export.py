"""
engine/excel_export.py
Excel Export Engine for Schedule III Finalisation.
Generates an executive, audit-grade multi-tab workbook (.xlsx):
- Sheet 1: Executive Summary & Audit Checks
- Sheet 2: Balance Sheet (CY & PY)
- Sheet 3: Statement of Profit & Loss (CY & PY)
- Sheet 4: Notes to Financial Statements (Ledger level schedules)
- Sheet 5: Mapped Trial Balance (Audit trail)
- Sheet 6: Schedule III Mandatory Ratios & 25% Variance Flags
- Sheet 7: CA Judgement & Assumptions Log
- Sheet 8: Statutory Disclosure Checklist (MCA 2021)
Uses openpyxl with clean typography, borders, and number formatting.
"""

import io
from typing import Dict, List, Any, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from engine.statements import GeneratedFinancialStatements
from engine.checks import AuditChecksSummary
from engine.ratios import RatioResult, ratios_to_dataframe
from engine.classifier import TreatmentLogEntry
from templates.schedule_iii_schema import MANDATORY_DISCLOSURE_CHECKLIST


def generate_schedule_iii_excel(
    stmts: GeneratedFinancialStatements,
    audit_summary: AuditChecksSummary,
    ratios: List[RatioResult],
    treatment_logs: List[TreatmentLogEntry],
    company_name: str = "Corporate Entity Limited"
) -> io.BytesIO:
    """
    Creates the comprehensive audit-ready Schedule III Excel workbook.
    Returns a BytesIO buffer suitable for download in Streamlit or writing to disk.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    navy_header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    soft_blue_fill = PatternFill(start_color="E8EEF5", end_color="E8EEF5", fill_type="solid")
    green_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
    amber_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    gray_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    font_title = Font(name="Calibri", size=15, bold=True, color="1B365D")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="595959")
    font_tbl_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_section_header = Font(name="Calibri", size=10, bold=True, color="1B365D")
    font_subtotal = Font(name="Calibri", size=10, bold=True, color="000000")
    font_total = Font(name="Calibri", size=11, bold=True, color="1B365D")
    font_regular = Font(name="Calibri", size=10, color="000000")
    font_bold = Font(name="Calibri", size=10, bold=True, color="000000")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    total_border = Border(
        top=Side(style="thin", color="000000"),
        bottom=Side(style="double", color="000000"),
    )

    # -------------------------------------------------------------
    # 1. EXECUTIVE SUMMARY & AUDIT CHECKS
    # -------------------------------------------------------------
    ws_sum = wb.create_sheet(title="Executive Summary & Checks")
    ws_sum.views.sheetView[0].showGridLines = True

    ws_sum.cell(row=2, column=2, value=f"{company_name}").font = font_title
    ws_sum.cell(row=3, column=2, value=f"Schedule III Finalisation Model | {stmts.division.value}").font = font_subtitle
    ws_sum.cell(row=4, column=2, value=f"Reporting Currency: INR | Rounding: {stmts.rounding_unit.value}").font = font_subtitle

    # Overall Status Panel
    status_text = "ALL STATUTORY CHECKS PASSED ✅" if audit_summary.all_passed else "ATTENTION: AUDIT CHECK EXCEPTIONS DETECTED ⚠️"
    ws_sum.cell(row=6, column=2, value="Finalisation Status:").font = font_bold
    cell_stat = ws_sum.cell(row=6, column=3, value=status_text)
    cell_stat.font = Font(name="Calibri", size=11, bold=True, color="276A3C" if audit_summary.all_passed else "C00000")
    cell_stat.fill = green_fill if audit_summary.all_passed else amber_fill

    headers_chk = ["Audit Check", "Status", "Difference (₹)", "Audit Reconciliation Summary"]
    for col_idx, h in enumerate(headers_chk, start=2):
        c = ws_sum.cell(row=8, column=col_idx, value=h)
        c.font = font_tbl_header
        c.fill = navy_header_fill
        c.alignment = Alignment(horizontal="center")

    curr_row = 9
    for chk in audit_summary.checks:
        ws_sum.cell(row=curr_row, column=2, value=chk.name).font = font_bold
        c_st = ws_sum.cell(row=curr_row, column=3, value=chk.status)
        c_st.font = Font(name="Calibri", size=10, bold=True, color="276A3C" if chk.is_passed else "C00000")
        c_st.alignment = Alignment(horizontal="center")
        c_st.fill = green_fill if chk.is_passed else amber_fill

        c_diff = ws_sum.cell(row=curr_row, column=4, value=chk.difference)
        c_diff.number_format = "#,##0.00"
        c_diff.font = font_regular
        c_diff.alignment = Alignment(horizontal="right")

        ws_sum.cell(row=curr_row, column=5, value=chk.summary_message).font = font_regular
        for col_idx in range(2, 6):
            ws_sum.cell(row=curr_row, column=col_idx).border = thin_border
        curr_row += 1

    # -------------------------------------------------------------
    # 2. BALANCE SHEET
    # -------------------------------------------------------------
    ws_bs = wb.create_sheet(title="Balance Sheet")
    ws_bs.views.sheetView[0].showGridLines = True

    ws_bs.cell(row=2, column=2, value=f"{company_name}").font = font_title
    ws_bs.cell(row=3, column=2, value=f"Balance Sheet as at 31st March (Schedule III {stmts.division.value})").font = font_subtitle
    ws_bs.cell(row=4, column=2, value=f"(Figures in {stmts.rounding_unit.value})").font = font_subtitle

    headers_bs = ["Particulars", "Note No.", f"As at 31st March, {stmts.cy_label[-4:]}", f"As at 31st March, {stmts.py_label[-4:]}"]
    for col_idx, h in enumerate(headers_bs, start=2):
        c = ws_bs.cell(row=6, column=col_idx, value=h)
        c.font = font_tbl_header
        c.fill = navy_header_fill
        c.alignment = Alignment(horizontal="left" if col_idx == 2 else "center")

    curr_row = 7
    for _, row in stmts.balance_sheet_df.iterrows():
        p = str(row.get("Particulars", ""))
        note = row.get("Note No.", "")
        cy_v = row.get(stmts.cy_label)
        py_v = row.get(stmts.py_label)
        is_hdr = bool(row.get("is_header", False))
        is_sub = bool(row.get("is_subtotal", False))
        is_tot = bool(row.get("is_total", False))

        c_p = ws_bs.cell(row=curr_row, column=2, value=p)
        c_note = ws_bs.cell(row=curr_row, column=3, value=note if note else "")
        c_note.alignment = Alignment(horizontal="center")
        c_cy = ws_bs.cell(row=curr_row, column=4, value=cy_v if pd.notna(cy_v) else "")
        c_py = ws_bs.cell(row=curr_row, column=5, value=py_v if pd.notna(py_v) else "")

        for col_idx in [2, 3, 4, 5]:
            ws_bs.cell(row=curr_row, column=col_idx).border = thin_border

        if is_tot:
            c_p.font = font_total
            c_cy.font = font_total
            c_py.font = font_total
            c_cy.fill = soft_blue_fill
            c_py.fill = soft_blue_fill
            c_cy.border = total_border
            c_py.border = total_border
        elif is_sub:
            c_p.font = font_subtotal
            c_cy.font = font_subtotal
            c_py.font = font_subtotal
            c_cy.fill = gray_fill
            c_py.fill = gray_fill
        elif is_hdr:
            c_p.font = font_section_header
        else:
            c_p.font = font_regular
            c_cy.font = font_regular
            c_py.font = font_regular

        if isinstance(cy_v, (int, float)):
            c_cy.number_format = "#,##0.00"
            c_cy.alignment = Alignment(horizontal="right")
        if isinstance(py_v, (int, float)):
            c_py.number_format = "#,##0.00"
            c_py.alignment = Alignment(horizontal="right")

        curr_row += 1

    # -------------------------------------------------------------
    # 3. STATEMENT OF PROFIT & LOSS
    # -------------------------------------------------------------
    ws_pl = wb.create_sheet(title="Profit & Loss")
    ws_pl.views.sheetView[0].showGridLines = True

    ws_pl.cell(row=2, column=2, value=f"{company_name}").font = font_title
    ws_pl.cell(row=3, column=2, value=f"Statement of Profit and Loss for the year ended 31st March (Schedule III {stmts.division.value})").font = font_subtitle
    ws_pl.cell(row=4, column=2, value=f"(Figures in {stmts.rounding_unit.value})").font = font_subtitle

    for col_idx, h in enumerate(headers_bs, start=2):
        c = ws_pl.cell(row=6, column=col_idx, value=h)
        c.font = font_tbl_header
        c.fill = navy_header_fill
        c.alignment = Alignment(horizontal="left" if col_idx == 2 else "center")

    curr_row = 7
    for _, row in stmts.pl_df.iterrows():
        p = str(row.get("Particulars", ""))
        note = row.get("Note No.", "")
        cy_v = row.get(stmts.cy_label)
        py_v = row.get(stmts.py_label)
        is_hdr = bool(row.get("is_header", False))
        is_sub = bool(row.get("is_subtotal", False))
        is_tot = bool(row.get("is_total", False))

        c_p = ws_pl.cell(row=curr_row, column=2, value=p)
        c_note = ws_pl.cell(row=curr_row, column=3, value=note if note else "")
        c_note.alignment = Alignment(horizontal="center")
        c_cy = ws_pl.cell(row=curr_row, column=4, value=cy_v if pd.notna(cy_v) else "")
        c_py = ws_pl.cell(row=curr_row, column=5, value=py_v if pd.notna(py_v) else "")

        for col_idx in [2, 3, 4, 5]:
            ws_pl.cell(row=curr_row, column=col_idx).border = thin_border

        if is_tot:
            c_p.font = font_total
            c_cy.font = font_total
            c_py.font = font_total
            c_cy.fill = soft_blue_fill
            c_py.fill = soft_blue_fill
            c_cy.border = total_border
            c_py.border = total_border
        elif is_sub:
            c_p.font = font_subtotal
            c_cy.font = font_subtotal
            c_py.font = font_subtotal
            c_cy.fill = gray_fill
            c_py.fill = gray_fill
        elif is_hdr:
            c_p.font = font_section_header
        else:
            c_p.font = font_regular
            c_cy.font = font_regular
            c_py.font = font_regular

        if isinstance(cy_v, (int, float)):
            c_cy.number_format = "#,##0.00"
            c_cy.alignment = Alignment(horizontal="right")
        if isinstance(py_v, (int, float)):
            c_py.number_format = "#,##0.00"
            c_py.alignment = Alignment(horizontal="right")

        curr_row += 1

    # -------------------------------------------------------------
    # 4. NOTES TO FINANCIAL STATEMENTS
    # -------------------------------------------------------------
    ws_notes = wb.create_sheet(title="Notes to Accounts")
    ws_notes.views.sheetView[0].showGridLines = True

    ws_notes.cell(row=2, column=2, value=f"{company_name}").font = font_title
    ws_notes.cell(row=3, column=2, value="Notes forming part of the Financial Statements (Ledger Schedules)").font = font_subtitle
    ws_notes.cell(row=4, column=2, value=f"(Figures in {stmts.rounding_unit.value})").font = font_subtitle

    curr_row = 6
    for note_title, note_obj in sorted(stmts.notes.items(), key=lambda x: x[1].note_no):
        if not note_obj.ledgers and note_obj.total_cy == 0 and note_obj.total_py == 0:
            continue

        ws_notes.cell(row=curr_row, column=2, value=note_obj.title).font = font_section_header
        ws_notes.cell(row=curr_row, column=2).fill = soft_blue_fill
        ws_notes.cell(row=curr_row, column=3).fill = soft_blue_fill
        ws_notes.cell(row=curr_row, column=4).fill = soft_blue_fill
        ws_notes.cell(row=curr_row, column=5).fill = soft_blue_fill
        curr_row += 1

        headers_note = ["Particulars / Ledger Account", "Group / Classification", f"{stmts.cy_label[-4:]} Amount", f"{stmts.py_label[-4:]} Amount"]
        for col_idx, h in enumerate(headers_note, start=2):
            c = ws_notes.cell(row=curr_row, column=col_idx, value=h)
            c.font = font_tbl_header
            c.fill = navy_header_fill
            c.alignment = Alignment(horizontal="left" if col_idx <= 3 else "right")
        curr_row += 1

        for ledger in note_obj.ledgers:
            c_p = ws_notes.cell(row=curr_row, column=2, value=ledger["ledger_name"])
            c_g = ws_notes.cell(row=curr_row, column=3, value=ledger.get("group", ""))
            c_cy = ws_notes.cell(row=curr_row, column=4, value=round(ledger["amount_cy"] / stmts.unit_multiplier, 4))
            c_py = ws_notes.cell(row=curr_row, column=5, value=round(ledger["amount_py"] / stmts.unit_multiplier, 4))

            c_p.font = font_regular
            c_g.font = font_regular
            c_cy.font = font_regular
            c_py.font = font_regular
            c_cy.number_format = "#,##0.00"
            c_py.number_format = "#,##0.00"
            c_cy.alignment = Alignment(horizontal="right")
            c_py.alignment = Alignment(horizontal="right")

            for col_idx in [2, 3, 4, 5]:
                ws_notes.cell(row=curr_row, column=col_idx).border = thin_border
            curr_row += 1

        # Total Row
        c_tot_lbl = ws_notes.cell(row=curr_row, column=2, value=f"Total: {note_obj.line_item_name}")
        c_tot_lbl.font = font_subtotal
        c_tot_cy = ws_notes.cell(row=curr_row, column=4, value=round(note_obj.total_cy / stmts.unit_multiplier, 4))
        c_tot_py = ws_notes.cell(row=curr_row, column=5, value=round(note_obj.total_py / stmts.unit_multiplier, 4))
        c_tot_cy.font = font_subtotal
        c_tot_py.font = font_subtotal
        c_tot_cy.number_format = "#,##0.00"
        c_tot_py.number_format = "#,##0.00"
        c_tot_cy.alignment = Alignment(horizontal="right")
        c_tot_py.alignment = Alignment(horizontal="right")

        for col_idx in [2, 3, 4, 5]:
            c_cell = ws_notes.cell(row=curr_row, column=col_idx)
            c_cell.fill = gray_fill
            c_cell.border = total_border
        curr_row += 2

    # -------------------------------------------------------------
    # 5. MAPPED TRIAL BALANCE (AUDIT TRAIL)
    # -------------------------------------------------------------
    ws_tb = wb.create_sheet(title="Mapped Trial Balance")
    ws_tb.views.sheetView[0].showGridLines = True

    ws_tb.cell(row=2, column=2, value=f"{company_name}").font = font_title
    ws_tb.cell(row=3, column=2, value="Mapped Trial Balance & Classification Workings").font = font_subtitle

    headers_tb = [
        "Ledger Name", "Group", "Debit (CY)", "Credit (CY)", "Net Balance (CY)",
        "Schedule III Target Line Item", "Statement", "Classification",
        "Confidence", "Source", "CA Rule Tag"
    ]
    for col_idx, h in enumerate(headers_tb, start=2):
        c = ws_tb.cell(row=5, column=col_idx, value=h)
        c.font = font_tbl_header
        c.fill = navy_header_fill
        c.alignment = Alignment(horizontal="center")

    curr_row = 6
    for _, row in stmts.drill_down_df.iterrows():
        ws_tb.cell(row=curr_row, column=2, value=str(row["Ledger Name"])).font = font_regular
        ws_tb.cell(row=curr_row, column=3, value=str(row["Group"])).font = font_regular

        c_dr = ws_tb.cell(row=curr_row, column=4, value=float(row["Debit (CY)"]))
        c_cr = ws_tb.cell(row=curr_row, column=5, value=float(row["Credit (CY)"]))
        c_net = ws_tb.cell(row=curr_row, column=6, value=float(row["Net Balance (CY)"]))
        for c in [c_dr, c_cr, c_net]:
            c.font = font_regular
            c.number_format = "#,##0.00"
            c.alignment = Alignment(horizontal="right")

        ws_tb.cell(row=curr_row, column=7, value=str(row["Schedule III Line Item"])).font = font_bold
        ws_tb.cell(row=curr_row, column=8, value=str(row["Statement"])).font = font_regular
        ws_tb.cell(row=curr_row, column=9, value=str(row["Classification"])).font = font_regular

        c_score = ws_tb.cell(row=curr_row, column=10, value=float(row["Match Score"]))
        c_score.font = font_regular
        c_score.number_format = "0.0%"
        c_score.alignment = Alignment(horizontal="center")

        ws_tb.cell(row=curr_row, column=11, value=str(row["Source"])).font = font_regular
        ws_tb.cell(row=curr_row, column=12, value=str(row["CA Rule Tag"])).font = font_regular

        for col_idx in range(2, 13):
            ws_tb.cell(row=curr_row, column=col_idx).border = thin_border
        curr_row += 1

    # -------------------------------------------------------------
    # 6. SCHEDULE III RATIOS & 25% VARIANCE ANALYSIS
    # -------------------------------------------------------------
    ws_ratios = wb.create_sheet(title="Schedule III Ratios")
    ws_ratios.views.sheetView[0].showGridLines = True

    ws_ratios.cell(row=2, column=2, value=f"{company_name}").font = font_title
    ws_ratios.cell(row=3, column=2, value="Mandatory 11 Schedule III Ratios Disclosure & Variance Flags (>25%)").font = font_subtitle

    headers_rat = [
        "Schedule III Mandatory Ratio", "Numerator ÷ Denominator",
        f"{stmts.cy_label[-4:]}", f"{stmts.py_label[-4:]}",
        "Variance %", "Flag (>25%)", "Statutory Note / CA Explanation"
    ]
    for col_idx, h in enumerate(headers_rat, start=2):
        c = ws_ratios.cell(row=5, column=col_idx, value=h)
        c.font = font_tbl_header
        c.fill = navy_header_fill
        c.alignment = Alignment(horizontal="center")

    curr_row = 6
    for r in ratios:
        ws_ratios.cell(row=curr_row, column=2, value=r.ratio_name).font = font_bold
        ws_ratios.cell(row=curr_row, column=3, value=f"{r.numerator_desc} ÷ {r.denominator_desc}").font = font_regular

        c_cy = ws_ratios.cell(row=curr_row, column=4, value=r.value_cy)
        c_py = ws_ratios.cell(row=curr_row, column=5, value=r.value_py)
        for c in [c_cy, c_py]:
            c.font = font_regular
            c.number_format = "#,##0.00" if r.unit_str != "%" else "0.00%"
            c.alignment = Alignment(horizontal="right")

        c_var = ws_ratios.cell(row=curr_row, column=6, value=f"{r.variance_pct:+.2f}%")
        c_var.font = font_bold
        c_var.alignment = Alignment(horizontal="right")

        c_flag = ws_ratios.cell(row=curr_row, column=7, value="⚠️ >25%" if r.is_variance_significant else "✅ OK")
        c_flag.font = Font(name="Calibri", size=10, bold=True, color="C00000" if r.is_variance_significant else "276A3C")
        c_flag.alignment = Alignment(horizontal="center")
        c_flag.fill = amber_fill if r.is_variance_significant else green_fill

        ws_ratios.cell(row=curr_row, column=8, value=r.ca_commentary).font = font_regular

        for col_idx in range(2, 9):
            ws_ratios.cell(row=curr_row, column=col_idx).border = thin_border
        curr_row += 1

    # -------------------------------------------------------------
    # 7. CA JUDGEMENT & ASSUMPTIONS LOG
    # -------------------------------------------------------------
    ws_log = wb.create_sheet(title="CA Judgement Log")
    ws_log.views.sheetView[0].showGridLines = True

    ws_log.cell(row=2, column=2, value=f"{company_name}").font = font_title
    ws_log.cell(row=3, column=2, value="Assumptions & Treatment Log (Reclassifications, Grossing Up & Netting)").font = font_subtitle

    headers_log = ["Ledger Account / Description", "Treatment Type", "Original Mapping", "Adjusted Presentation", "Amount (₹)", "Statutory Rationale"]
    for col_idx, h in enumerate(headers_log, start=2):
        c = ws_log.cell(row=5, column=col_idx, value=h)
        c.font = font_tbl_header
        c.fill = navy_header_fill
        c.alignment = Alignment(horizontal="center")

    curr_row = 6
    for item in treatment_logs:
        ws_log.cell(row=curr_row, column=2, value=item.ledger_name).font = font_bold
        ws_log.cell(row=curr_row, column=3, value=item.treatment_type).font = font_regular
        ws_log.cell(row=curr_row, column=4, value=item.original_line_item).font = font_regular
        ws_log.cell(row=curr_row, column=5, value=item.adjusted_line_item).font = font_bold

        c_amt = ws_log.cell(row=curr_row, column=6, value=item.amount_cy)
        c_amt.font = font_regular
        c_amt.number_format = "#,##0.00"
        c_amt.alignment = Alignment(horizontal="right")

        ws_log.cell(row=curr_row, column=7, value=item.statutory_rationale).font = font_regular

        for col_idx in range(2, 8):
            ws_log.cell(row=curr_row, column=col_idx).border = thin_border
        curr_row += 1

    # -------------------------------------------------------------
    # 8. STATUTORY DISCLOSURE CHECKLIST
    # -------------------------------------------------------------
    ws_disc = wb.create_sheet(title="Disclosure Checklist")
    ws_disc.views.sheetView[0].showGridLines = True

    ws_disc.cell(row=2, column=2, value=f"{company_name}").font = font_title
    ws_disc.cell(row=3, column=2, value="Schedule III Non-TB Additional Statutory Disclosures (MCA 2021 Checklist)").font = font_subtitle

    headers_disc = ["Statutory Disclosure Item", "Companies Act / MCA Reference", "Scope & Required Details", "Audit Status", "Input Fields Required"]
    for col_idx, h in enumerate(headers_disc, start=2):
        c = ws_disc.cell(row=5, column=col_idx, value=h)
        c.font = font_tbl_header
        c.fill = navy_header_fill
        c.alignment = Alignment(horizontal="center")

    curr_row = 6
    for d in MANDATORY_DISCLOSURE_CHECKLIST:
        ws_disc.cell(row=curr_row, column=2, value=d["item"]).font = font_bold
        ws_disc.cell(row=curr_row, column=3, value=d["reference"]).font = font_regular
        ws_disc.cell(row=curr_row, column=4, value=d["description"]).font = font_regular

        c_st = ws_disc.cell(row=curr_row, column=5, value=d["status"])
        c_st.font = font_bold
        c_st.alignment = Alignment(horizontal="center")
        c_st.fill = amber_fill if "Input" in d["status"] else green_fill

        ws_disc.cell(row=curr_row, column=6, value=", ".join(d["fields"])).font = font_regular

        for col_idx in range(2, 7):
            ws_disc.cell(row=curr_row, column=col_idx).border = thin_border
        curr_row += 1

    # Auto-adjust column widths across all sheets
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val:
                    val_str = str(val)
                    if len(val_str) > max_len and len(val_str) < 70:
                        max_len = len(val_str)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
